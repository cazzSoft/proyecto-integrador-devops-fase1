from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Proyecto_Integrador_DevOps_Fase_1.docx"
LOGO = ROOT / "docs" / "assets" / "logoespam.png"

NAVY = "12304A"
GREEN = "006B58"
DARK_GREEN = "00483D"
BLUE = "147EA3"
LIGHT_GREEN = "E9F5F1"
LIGHT_BLUE = "EAF4F8"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "667785"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def set_font(run, size=10, color="263844", bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_para(doc, text="", size=10, color="263844", bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=5, line=1.08):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_font(r, size, color, bold, italic)
    return p


def add_heading(doc, number, title, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    r = p.add_run(f"{number}  {title}" if number else title)
    set_font(r, 16 if level == 1 else 11.5, NAVY if level == 1 else GREEN, True)
    return p


def add_kicker(doc, text):
    return add_para(doc, text.upper(), size=8, color=GREEN, bold=True,
                    align=WD_ALIGN_PARAGRAPH.LEFT, after=2, line=1.0)


def add_table(doc, headers, rows, widths, font_size=7.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(text), font_size, WHITE, True)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            set_cell_shading(cell, WHITE if row_idx % 2 == 0 else LIGHT_GRAY)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(str(text)), font_size, "334954")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label, text, fill=LIGHT_GREEN):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label), 8.5, DARK_GREEN, True)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(text), 8.5, "334954")
    return table


def add_flow(doc, labels, caption):
    table = doc.add_table(rows=1, cols=len(labels))
    table.style = "Table Grid"
    widths = [9360 // len(labels)] * len(labels)
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    for idx, label in enumerate(labels):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GREEN if idx < len(labels) - 2 else "FFF2D9")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(f"{idx + 1}\n{label}"), 7.1, NAVY, True)
    add_para(doc, caption, size=7.2, color=MID_GRAY, italic=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4, line=1.0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, 8, MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)
section.different_first_page_header_footer = True

# Preset: compact_reference_guide. Named override: academic_density (10 pt body,
# 0.78 in side margins and 7.1-8 pt table text) to honor the four-page content limit.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string("263844")
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for style_name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 11.5, GREEN), ("Heading 3", 10.5, DARK_GREEN)):
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)

header = section.header
p = header.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(p.add_run("ESPAM MFL  |  DEVOPS  |  PROYECTO INTEGRADOR - FASE 1"), 8, NAVY, True)
footer = section.footer
add_page_number(footer.paragraphs[0])
first_footer = section.first_page_footer
fp = first_footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("ESPAM MFL | Maestría en Diseño Web y Desarrollo de Apps | Página 1"), 7.5, MID_GRAY)

# Portada
add_para(doc, "", after=6)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo_shape = p.add_run().add_picture(str(LOGO), width=Inches(3.75))
logo_shape._inline.docPr.set("descr", "Logotipo institucional de la ESPAM MFL")
logo_shape._inline.docPr.set("title", "ESPAM MFL")
add_para(doc, "", after=7)
add_para(doc, "ESCUELA SUPERIOR POLITÉCNICA AGROPECUARIA DE MANABÍ\nMANUEL FÉLIX LÓPEZ",
         size=11, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.0)
add_para(doc, "MAESTRÍA EN DISEÑO WEB Y DESARROLLO DE APPS\nII COHORTE PARALELO A",
         size=10, color="263844", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12, line=1.0)
add_para(doc, "TEMA:", size=10, color="263844", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.0)
add_para(doc, "PROYECTO INTEGRADOR DEVOPS - FASE 1",
         size=18, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.0)
add_para(doc, "DIAGNÓSTICO Y PLANIFICACIÓN",
         size=12, color=GREEN, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.0)
add_para(doc, "SISTEMA SELECCIONADO:", size=9, color="263844", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.0)
add_para(doc, "SISTEMA WEB CRUD PARA LA GESTIÓN DE USUARIOS",
         size=9.5, color="263844", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=9, line=1.0)
add_para(doc, "AUTOR:", size=9, color="263844", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.0)
add_para(doc, "CRISTHIAN ALFREDO ZAMBRANO ZAMBRANO",
         size=9.5, color="263844", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.0)
add_para(doc, "MÓDULO: DEVOPS", size=9, color="263844", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.0)
add_para(doc, "DOCENTE: MG. ENRIQUE JAVIER MACÍAS ARIAS",
         size=9, color="263844", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=9, line=1.0)
add_para(doc, "AGOSTO DE 2026", size=9.5, color="263844", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
doc.add_page_break()

# Página 1 de contenido
add_kicker(doc, "01 · Contexto y flujo")
add_heading(doc, "", "Sistema web CRUD para la gestión de usuarios")
add_para(doc, "Aplicación académica que centraliza la creación, consulta, actualización y eliminación de usuarios, además de asignar roles. Beneficia a administradores y personal de soporte que requieren mantener registros consistentes mediante una interfaz web.")
add_table(doc, ["Capa", "Tecnologías", "Responsabilidad"], [
    ["Frontend", "React 18, TypeScript, Vite y Mantine", "Interfaz, formularios, tabla y consumo de la API REST."],
    ["Backend", "Node.js, Fastify, TypeScript y Prisma", "Reglas del CRUD, validación básica y acceso a datos."],
    ["Datos", "PostgreSQL 17", "Persistencia de usuarios."],
    ["Ejecución", "Docker y Docker Compose", "Construcción y coordinación de tres servicios."],
], [1500, 3500, 4360])
add_heading(doc, "1.1", "Arquitectura actual", 2)
add_flow(doc, ["Administrador", "Frontend\nReact", "API\nFastify", "Prisma", "PostgreSQL"],
         "Figura 1. Flujo de solicitudes en la arquitectura actual.")
add_heading(doc, "1.2", "Flujo actual del código al despliegue", 2)
add_para(doc, "La línea base corresponde a la práctica recibida. No existe CI/CD: una persona modifica, construye, valida y publica; las esperas no están medidas y la evidencia se limita a la consola y a la comprobación visual.", size=9.2)
add_flow(doc, ["Modificar\ncódigo", "Guardar\ncambios", "Compose\nbuild", "Instalar\ndependencias", "Prisma\ndb push", "Probar\nmanual", "Publicar\nmanual"],
         "Figura 2. Si la comprobación falla, el desarrollador corrige y repite manualmente el flujo.")
add_table(doc, ["Etapa", "Responsable/evidencia", "Actividad manual o riesgo"], [
    ["Código e integración", "Desarrollador; archivos y commit", "Revisión sin validación automática."],
    ["Build y datos", "Contenedores; salida de consola", "Dependencias y db push al arrancar."],
    ["Prueba y aprobación", "Desarrollador; navegador", "CRUD manual, sin suite ni criterio independiente."],
    ["Despliegue/operación", "Desarrollador; consola y logs", "Sin pipeline, rollback, alerta o medición."],
], [2500, 3100, 3760])
doc.add_page_break()

# Página 2 de contenido
add_kicker(doc, "02 · Diagnóstico")
add_heading(doc, "", "Cuellos de botella, actividades manuales y riesgos")
add_para(doc, "La prioridad combina impacto y probabilidad de 1 (bajo) a 3 (alto). La puntuación es su producto; los valores 9 se atienden primero.", size=9.2)
add_table(doc, ["Hallazgo/tipo", "Causa y evidencia", "Consecuencia", "I×P"], [
    ["Sin CI/CD\nProceso/herramienta", "Toda validación depende de comandos locales.", "Errores tardíos y baja trazabilidad.", "3×3=9"],
    ["Sin pruebas\nProceso", "No existen suites unitarias, API o e2e.", "Regresiones, retrabajo y aprobación subjetiva.", "3×3=9"],
    ["Dependencias vulnerables\nSeguridad", "npm audit: 2 hallazgos frontend y 6 altos backend.", "Riesgo de cadena de suministro.", "3×3=9"],
    ["Credenciales originales\nSeguridad", "Contraseña en .env y Compose; retirada del repositorio.", "Exposición y acceso no autorizado.", "3×2=6"],
    ["Sin observabilidad\nOperación", "Logs reactivos, sin métricas o alertas.", "Detección y diagnóstico tardíos.", "2×3=6"],
    ["Sin rollback/respaldo\nContinuidad", "No existe runbook o restauración probada.", "Mayor indisponibilidad o pérdida de datos.", "3×2=6"],
], [2400, 3100, 3000, 860], 7.2)
add_heading(doc, "2.1", "Diagnóstico CALMS", 2)
add_table(doc, ["Dimensión", "Estado actual", "Brecha y respuesta"], [
    ["Cultura", "Responsabilidad concentrada en una persona.", "Compartir decisiones y revisar fallos sin culpa."],
    ["Automatización", "Compose reproduce los servicios.", "Automatizar build, pruebas, seguridad y entrega."],
    ["Lean", "Arquitectura simple.", "Reducir lotes y retrabajo con validación temprana."],
    ["Medición", "Logs de Fastify, sin línea base.", "Recopilar DORA, cobertura e incidentes."],
    ["Sharing", "Código recibido como ZIP.", "Repositorio, README, diagramas y runbook."],
], [1700, 3400, 4260], 7.4)
add_callout(doc, "Prioridad inmediata", "Automatizar compilación y pruebas, bloquear secretos y visibilizar dependencias vulnerables antes de automatizar el despliegue.")
doc.add_page_break()

# Página 3 de contenido
add_kicker(doc, "03 · Propuesta inicial")
add_heading(doc, "", "Adopción DevOps incremental y medible")
add_callout(doc, "Objetivo", "Reducir errores manuales y hacer reproducible, segura y medible la entrega mediante validaciones tempranas, configuración protegida y retroalimentación operativa.")
add_heading(doc, "3.1", "Alcance y arquitectura preliminar", 2)
add_table(doc, ["Dentro del alcance", "Fuera del alcance inicial"], [[
    "GitHub; commits y revisión; CI; pruebas; SAST; auditoría de dependencias/secretos; imágenes Docker; QA; observabilidad y rollback.",
    "Kubernetes, alta disponibilidad multirregión y despliegue automático a producción: añadirían complejidad antes de estabilizar el flujo.",
]], [4900, 4460], 7.8)
add_flow(doc, ["Push", "Build", "Pruebas", "Seguridad", "Imagen", "QA", "Desplegar", "Observar"],
         "Figura 3. Flujo propuesto con retroalimentación desde operación.")
add_heading(doc, "3.2", "Trazabilidad problema-práctica-métrica", 2)
add_table(doc, ["Problema", "Práctica/herramienta", "Métrica", "Resultado"], [
    ["Validación manual", "CI con GitHub Actions o Jenkins.", "Lead time/frecuencia", "Evidencia automática."],
    ["Pruebas inexistentes", "Pruebas unitarias y API.", "Cobertura/fallos", "Menos regresión."],
    ["Dependencias/secretos", "npm audit, Dependabot, SAST y Secrets.", "Hallazgos/versión", "Riesgo visible temprano."],
    ["Detección tardía", "Logs, métricas y OpenTelemetry.", "Detección/MTTR", "Respuesta más rápida."],
    ["Recuperación incierta", "Imágenes, rollback, respaldo y runbook.", "MTTR/restauración", "Continuidad verificable."],
], [2100, 3000, 2000, 2260], 7.2)
add_heading(doc, "3.3", "Roadmap de 90 días", 2)
add_table(doc, ["Periodo", "Cultura", "Proceso", "Técnica/evidencia"], [
    ["0-30", "Responsabilidad compartida.", "Commits, revisión y criterios mínimos.", "CI de build + secretos; resultado por cambio."],
    ["31-60", "Revisión sin culpa.", "Aceptación y rollback.", "Pruebas, SAST y dependencias visibles."],
    ["61-90", "Revisión de métricas.", "Cambios pequeños y aprobación QA.", "Imágenes, observabilidad y runbook."],
], [1200, 2200, 2700, 3260], 7.2)
doc.add_page_break()

# Página 4 de contenido
add_kicker(doc, "04 · Medición y cierre")
add_heading(doc, "", "Línea base, metas y gestión de riesgos")
add_para(doc, "No se inventan cifras históricas: la ausencia de medición es parte del diagnóstico. En los primeros 30 días el pipeline y el registro de incidentes generarán una línea base reproducible.", size=9.2)
add_table(doc, ["Métrica", "Base", "Obtención", "Meta a 90 días"], [
    ["Lead time", "No medido", "Commit a despliegue según CI/CD.", "Línea base y tendencia descendente."],
    ["Frecuencia", "No registrada", "Despliegues exitosos por mes.", "Al menos 2 controlados/mes."],
    ["Fallos del cambio", "No medida", "Rollback o corrección / total.", "Menor al 20 % tras la base."],
    ["Recuperación", "No medido", "Incidente a restauración.", "Runbook y menos de 60 min."],
    ["Cobertura", "0 % / sin suite", "Reporte automático del pipeline.", "60 % en lógica crítica."],
], [2100, 1600, 2900, 2760], 7.4)
add_heading(doc, "4.1", "Riesgos del cambio y controles", 2)
add_table(doc, ["Riesgo", "Control y evidencia"], [
    ["Seguridad/privacidad", ".env excluido, secretos protegidos, escaneo por commit, minimización y retención de logs."],
    ["Continuidad", "Healthchecks, respaldo/restauración, imágenes versionadas, rollback y runbook probado."],
    ["Adopción", "Automatización gradual, guía de contribución y revisión mensual de métricas."],
], [2500, 6860], 7.8)
add_heading(doc, "4.2", "Conclusiones", 2)
add_para(doc, "El sistema es apropiado porque integra interfaz, API y base de datos, pero su flujo original depende de validaciones y despliegues manuales. Los cuellos críticos deben resolverse primero con compilación, pruebas y seguridad tempranas; después, entrega controlada, observabilidad y recuperación. La propuesta conecta cada problema con una práctica, evidencia y métrica, por lo que permite evaluar mejoras sin comparar personas ni inventar datos.", size=9.2)
add_heading(doc, "", "Referencias (APA 7)", 2)
refs = [
    "DORA. (2024). Accelerate State of DevOps Report 2024. Google Cloud.",
    "ESPAM MFL. (2026). Guía de estudio: Unidad 1 DevOps. Maestría en Diseño Web y Desarrollo de Apps.",
    "ESPAM MFL. (2026). Presentación DevOps, Unidad 1, Sesiones 1-3 [Diapositivas de clase].",
    "Kim, G., Humble, J., Debois, P., Willis, J., & Forsgren, N. (2021). The DevOps handbook (2nd ed.). IT Revolution.",
]
for ref in refs:
    p = add_para(doc, ref, size=7.6, color="334954", align=WD_ALIGN_PARAGRAPH.LEFT, after=2, line=1.0)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
add_callout(doc, "Repositorio público", "https://github.com/cazzSoft/proyecto-integrador-devops-fase1", LIGHT_BLUE)

doc.core_properties.title = "Proyecto integrador DevOps - Fase 1"
doc.core_properties.subject = "Diagnóstico y planificación DevOps"
doc.core_properties.author = "Cristhian Alfredo Zambrano Zambrano"
doc.core_properties.keywords = "DevOps, diagnóstico, planificación, CALMS, DORA"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
